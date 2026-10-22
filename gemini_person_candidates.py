"""Use Gemini to attach player/team sentiment labels to saved post titles.

The original CSV remains untouched. One post can produce several rows when it
mentions several players. The local nickname workbook is supplied to Gemini as
reference material in every request.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import time
from pathlib import Path

try:
    from google import genai
    from google.genai import errors, types
    from openpyxl import load_workbook
except ImportError:
    print("필요 패키지가 없습니다. 아래 명령을 먼저 실행하세요:")
    print('  & "C:\\Users\\USER\\AppData\\Local\\Python\\pythoncore-3.14-64\\python.exe" -m pip install -U google-genai openpyxl')
    raise SystemExit(1)


BASE_DIR = Path(__file__).resolve().parent
INPUT_CSV = BASE_DIR / "dcinside_basketballkorea_posts.csv"
ALIAS_XLSX = BASE_DIR / "KBL_선수_별명 총정리.xlsx"
OUTPUT_CSV = BASE_DIR / "dcinside_basketballkorea_posts_선수평가.csv"
CHECKPOINT_FILE = BASE_DIR / "gemini_post_evaluation_checkpoint.json"

# Speed/cost-first default for the large title batches.  A caller can override
# either model through environment variables without changing source.
MODEL = os.getenv("GEMINI_MODEL", "gemini-flash-lite-latest")
FALLBACK_MODEL = os.getenv("GEMINI_FALLBACK_MODEL", "gemini-flash-latest")
BATCH_SIZE = 100
PAUSE_SECONDS = 1.0
MAX_SERVER_RETRIES = 2

HOME_TEAM = "울산 현대모비스"
AWAY_TEAM = "원주 DB"
# 기본값은 수동 실행 호환용이다. 자동 실행에서는 KBL API 명단으로 덮어쓴다.
PLAYED_PLAYERS = {
    "정준원", "이승현", "서명진", "김경원", "이우석", "게이지 프림", "김태완", "박무빈",
    "다리우스 베즐리", "제리 아바디아노", "이정현", "강상재", "최성원", "이윤기", "이용우",
    "이선 알바노", "레이션 해먼즈", "헨리 엘런슨", "이유진", "송재환",
}
PLAYED_ROSTER: dict[str, str] = {player: "" for player in PLAYED_PLAYERS}
GAME_DATE = os.getenv("KBL_GAME_DATE", "")


def clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def read_posts() -> list[dict[str, str]]:
    if not INPUT_CSV.exists():
        raise FileNotFoundError(f"입력 파일이 없습니다: {INPUT_CSV}")
    with INPUT_CSV.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)
        if not reader.fieldnames or "제목" not in reader.fieldnames:
            raise ValueError("입력 CSV에 '제목' 열이 없습니다.")
        return list(reader)


def load_alias_reference() -> list[tuple[str, str]]:
    """Read the user-maintained nickname dictionary without altering it."""
    if not ALIAS_XLSX.exists():
        raise FileNotFoundError(f"별명 참고 파일이 없습니다: {ALIAS_XLSX}")
    workbook = load_workbook(ALIAS_XLSX, read_only=True, data_only=True)
    sheet = workbook.active
    aliases: list[tuple[str, str]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        alias, explanation = (list(row) + [None, None])[:2]
        alias_text, explanation_text = clean_text(alias), clean_text(explanation)
        if alias_text and explanation_text:
            aliases.append((alias_text, explanation_text))
    if not aliases:
        raise ValueError("별명 참고 파일에 별명/해설 데이터가 없습니다.")
    return aliases


def alias_reference_text(aliases: list[tuple[str, str]]) -> str:
    return "\n".join(f"- {alias} → {explanation}" for alias, explanation in aliases)


def roster_reference_text() -> str:
    return "\n".join(
        f"- {team or '소속팀 확인'}: {player}"
        for player, team in sorted(PLAYED_ROSTER.items(), key=lambda item: (item[1], item[0]))
    )


def make_prompt(posts: list[dict[str, str]], post_ids: list[int], aliases: list[tuple[str, str]]) -> str:
    numbered_titles = "\n".join(
        f"{post_id}\t{post.get('제목', '').strip()}"
        for post_id, post in zip(post_ids, posts)
    )
    return f"""역할: KBL 커뮤니티 제목을 구조화하는 데이터 라벨러다.
대상 경기: {GAME_DATE or '당일'} {HOME_TEAM} 대 {AWAY_TEAM}. 이 두 팀/선수에 대한 글만 판독한다.
KBL 공식 기록으로 확정된 실제 출전 선수(아래 명단 밖의 선수 이름은 절대 선수로 반환하지 않는다):
{roster_reference_text()}
별명 사전(반드시 참고):
{alias_reference_text(aliases)}

반드시 지킬 출력 규칙:
1. 입력의 모든 id에 대해 JSON 객체를 최소 1개 반환한다. 설명·마크다운·코드블록은 절대 쓰지 말고 JSON 배열만 반환한다.
2. 한 제목에 두 선수 이상이 있으면 선수별로 객체를 각각 반환한다.
3. 선수 객체는 반드시 위 실제 출전 명단의 실명만 반환한다. 별명·오타·축약어를 명단 내 실명으로 확실히 연결할 수 있을 때만 선수로 반환한다.
4. 명단 밖 선수(벤치 미출전, 이전 선수, 타 팀 선수, 코치·인물 포함)를 가리키거나 선수가 불확실하면, 그 선수로 추정하지 말고 관련 팀의 팀 전체로 반환한다. id를 생략하지 않는다.
4. 팀 전체 글은 대상종류를 정확히 \"팀 전체\"로, 추정대상을 정확히 \"팀 전체: {HOME_TEAM}\" 또는 \"팀 전체: {AWAY_TEAM}\"로 쓴다.
5. 선수 글은 대상종류를 정확히 \"선수\"로, 추정대상은 위 명단에 있는 실명으로 쓴다.
6. 소속팀은 정확히 \"{HOME_TEAM}\" 또는 \"{AWAY_TEAM}\" 둘 중 하나만 쓴다.
7. 평가점수는 -5부터 5까지의 정수: 칭찬 양수, 비판 음수, 뉴스/질문/잡담 0.

반환 객체 키는 정확히 id, 추정대상, 대상종류, 소속팀, 평가점수 다섯 개다.

제목(ID<TAB>제목):
{numbered_titles}"""


def normalize_team(value: object) -> str:
    raw = clean_text(value)
    for team in (HOME_TEAM, AWAY_TEAM):
        short = team.replace("안양 ", "").replace("고양 ", "").replace("울산 ", "").replace("원주 ", "")
        if raw == team or raw == short or raw in team or short in raw:
            return team
    return ""


def normalize_item(item: dict[str, object], expected_ids: set[int]) -> tuple[int, dict[str, object]] | None:
    # Flash 계열은 드물게 키 앞에 공백을 붙인다(예: " 평가점수").
    # 의미는 동일하므로 키 공백만 정리하고 값은 그대로 보존한다.
    item = {clean_text(key): value for key, value in item.items()}
    try:
        post_id = int(item.get("id"))
        score = int(item.get("평가점수"))
    except (TypeError, ValueError):
        return None
    target = clean_text(item.get("추정대상"))
    kind = clean_text(item.get("대상종류"))
    team = normalize_team(item.get("소속팀"))
    if post_id not in expected_ids or not target or not -5 <= score <= 5:
        return None
    if kind.replace(" ", "") in {"팀", "팀전체", "team", "teamoverall"}:
        if not team:
            team = normalize_team(target)
        if not team:
            return None
        kind = "팀 전체"
        target = f"팀 전체: {team}"
    elif "선수" in kind or kind.lower() in {"player", "person"}:
        if not team:
            return None
        if target not in PLAYED_PLAYERS:
            # Never let a former player, non-participant or unrelated person
            # contaminate a player's score. Keep the sentiment at team level.
            kind = "팀 전체"
            target = f"팀 전체: {team}"
        else:
            kind = "선수"
    else:
        return None
    direction = "긍정" if score > 0 else "부정" if score < 0 else "중립"
    return post_id, {
        "추정대상": target,
        "대상종류": kind,
        "소속팀": team,
        "출전여부": "출전" if kind == "선수" else "해당 없음",
        "평가점수": score,
        "평가방향": direction,
    }


def parse_or_mark_missing(text: str, expected_ids: set[int], label: str) -> tuple[dict[int, list[dict[str, object]]], set[int]]:
    try:
        cleaned = text.strip()
        if cleaned.startswith("```"):
            cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE)
        data = json.loads(cleaned)
        if not isinstance(data, list):
            raise ValueError("JSON 배열이 아닙니다")
    except (json.JSONDecodeError, ValueError) as error:
        print(f"{label}: Gemini JSON 형식 오류로 재시도 대상 처리: {error}")
        return {}, expected_ids

    grouped: dict[int, list[dict[str, object]]] = {}
    seen: set[tuple[int, str]] = set()
    for raw_item in data:
        if not isinstance(raw_item, dict):
            continue
        normalized = normalize_item(raw_item, expected_ids)
        if not normalized:
            continue
        post_id, item = normalized
        key = (post_id, str(item["추정대상"]))
        if key not in seen:
            grouped.setdefault(post_id, []).append(item)
            seen.add(key)
    return grouped, expected_ids - grouped.keys()


def fallback_annotation(title: str) -> dict[str, object]:
    team = AWAY_TEAM if "db" in title.lower() or "원주" in title else HOME_TEAM
    return {
        "추정대상": f"팀 전체: {team}", "대상종류": "팀 전체", "소속팀": team,
        "출전여부": "해당 없음", "평가점수": 0, "평가방향": "중립",
    }


def load_checkpoint(post_count: int) -> tuple[int, dict[int, list[dict[str, object]]]]:
    if not CHECKPOINT_FILE.exists():
        return 0, {}
    try:
        saved = json.loads(CHECKPOINT_FILE.read_text(encoding="utf-8"))
        valid = (
            saved.get("kind") == "post_evaluation_multiple_v4"
            and saved.get("input_file") == INPUT_CSV.name
            and saved.get("post_count") == post_count
        )
        if not valid:
            print("이전 방식의 체크포인트는 사용하지 않고 처음부터 시작합니다.")
            return 0, {}
        return int(saved["completed_batches"]), {int(key): value for key, value in saved["annotations"].items()}
    except (OSError, KeyError, TypeError, ValueError):
        print("체크포인트를 읽지 못해 처음부터 시작합니다.")
        return 0, {}


def save_checkpoint(post_count: int, completed_batches: int, annotations: dict[int, list[dict[str, object]]]) -> None:
    payload = {
        "kind": "post_evaluation_multiple_v4", "input_file": INPUT_CSV.name,
        "post_count": post_count, "batch_size": BATCH_SIZE,
        "completed_batches": completed_batches, "annotations": annotations,
    }
    CHECKPOINT_FILE.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def call_gemini(client: genai.Client, prompt: str, batch_label: str):
    models = [MODEL] + ([FALLBACK_MODEL] if FALLBACK_MODEL and FALLBACK_MODEL != MODEL else [])
    last_error: Exception | None = None
    for model in models:
        for attempt in range(1, MAX_SERVER_RETRIES + 1):
            try:
                return client.models.generate_content(
                    model=model, contents=prompt,
                    config=types.GenerateContentConfig(response_mime_type="application/json", temperature=0.0, max_output_tokens=8192),
                )
            except errors.ServerError as error:
                last_error = error
                if attempt == MAX_SERVER_RETRIES:
                    break
                seconds = 3 * (2 ** (attempt - 1))
                print(f"{batch_label}: Gemini 503. {seconds}초 후 재시도 ({attempt}/{MAX_SERVER_RETRIES})")
                time.sleep(seconds)
    raise RuntimeError("Gemini 서버 과부하로 배치를 처리하지 못했습니다. 체크포인트는 보존됩니다.") from last_error


def evaluate_batch(client: genai.Client, posts: list[dict[str, str]], post_ids: list[int], aliases: list[tuple[str, str]], batch_number: int) -> dict[int, list[dict[str, object]]]:
    expected = set(post_ids)
    response = call_gemini(client, make_prompt(posts, post_ids, aliases), str(batch_number))
    raw = response.text or ""
    debug_file = CHECKPOINT_FILE.parent / f"gemini_raw_batch_{batch_number:03d}.txt"
    debug_file.write_text(raw, encoding="utf-8")
    result, missing = parse_or_mark_missing(raw, expected, f"{batch_number}번째 배치")
    if len(missing) == len(expected):
        print(f"{batch_number}번째 배치: 유효 응답 0개. 원문 저장: {debug_file}")
    for retry in range(1, 3):
        if not missing:
            break
        ids = sorted(missing)
        subset = [posts[post_ids.index(post_id)] for post_id in ids]
        print(f"{batch_number}번째 배치: 누락 {len(ids)}개만 재요청 ({retry}/2)")
        response = call_gemini(client, make_prompt(subset, ids, aliases), f"{batch_number}-누락{retry}")
        raw = response.text or ""
        (CHECKPOINT_FILE.parent / f"gemini_raw_batch_{batch_number:03d}_retry{retry}.txt").write_text(raw, encoding="utf-8")
        retried, missing = parse_or_mark_missing(raw, set(ids), f"{batch_number}번째 배치 재요청 {retry}")
        result.update(retried)
    for post_id in missing:
        result[post_id] = [fallback_annotation(posts[post_ids.index(post_id)].get("제목", ""))]
    return result


def write_output(posts: list[dict[str, str]], annotations: dict[int, list[dict[str, object]]]) -> None:
    original_columns = list(posts[0].keys()) if posts else []
    added = ["추정대상", "대상종류", "소속팀", "출전여부", "평가점수", "평가방향"]
    with OUTPUT_CSV.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=original_columns + added, extrasaction="ignore")
        writer.writeheader()
        for post_id, post in enumerate(posts, start=1):
            for annotation in annotations[post_id]:
                writer.writerow({**post, **annotation})


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Gemini로 DCInside 제목의 KBL 선수/팀 평가를 분류합니다.")
    parser.add_argument("--input", type=Path, default=INPUT_CSV)
    parser.add_argument("--output", type=Path, default=OUTPUT_CSV)
    parser.add_argument("--checkpoint", type=Path, default=CHECKPOINT_FILE)
    parser.add_argument("--aliases", type=Path, default=ALIAS_XLSX)
    parser.add_argument("--home", default=HOME_TEAM)
    parser.add_argument("--away", default=AWAY_TEAM)
    parser.add_argument("--played", default="", help="실제 출전 선수 이름을 | 로 구분")
    parser.add_argument("--roster-json", default="", help="KBL API가 확정한 [{team, player}] 실제 출전 명단 JSON")
    parser.add_argument("--date", default=GAME_DATE, help="프롬프트에 표시할 YYYYMMDD")
    parser.add_argument("--batch-size", type=int, default=BATCH_SIZE)
    return parser.parse_args()


def main() -> None:
    global INPUT_CSV, OUTPUT_CSV, CHECKPOINT_FILE, ALIAS_XLSX, HOME_TEAM, AWAY_TEAM, PLAYED_PLAYERS, PLAYED_ROSTER, BATCH_SIZE, GAME_DATE
    args = parse_args()
    INPUT_CSV, OUTPUT_CSV = args.input.resolve(), args.output.resolve()
    CHECKPOINT_FILE, ALIAS_XLSX = args.checkpoint.resolve(), args.aliases.resolve()
    HOME_TEAM, AWAY_TEAM, BATCH_SIZE, GAME_DATE = args.home, args.away, args.batch_size, args.date
    if args.roster_json:
        try:
            raw_roster = json.loads(args.roster_json)
            PLAYED_ROSTER = {
                clean_text(item.get("player")): clean_text(item.get("team"))
                for item in raw_roster if isinstance(item, dict) and clean_text(item.get("player"))
            }
        except json.JSONDecodeError as error:
            raise ValueError("--roster-json은 유효한 JSON 배열이어야 합니다.") from error
        PLAYED_PLAYERS = set(PLAYED_ROSTER)
    elif args.played:
        PLAYED_PLAYERS = {clean_text(name) for name in args.played.split("|") if clean_text(name)}
        PLAYED_ROSTER = {name: "" for name in PLAYED_PLAYERS}
    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    if not os.environ.get("GEMINI_API_KEY"):
        raise RuntimeError('GEMINI_API_KEY 환경변수가 비어 있습니다. 예: $env:GEMINI_API_KEY = "본인의_API_키"')
    posts, aliases = read_posts(), load_alias_reference()
    completed, annotations = load_checkpoint(len(posts))
    batches = [posts[start : start + BATCH_SIZE] for start in range(0, len(posts), BATCH_SIZE)]
    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
    if annotations:
        print(f"원본 {len(posts)}개 글과 별명 {len(aliases)}개를 읽었습니다. {len(annotations)}개 글의 기존 결과를 이어서 사용합니다.")
    else:
        print(f"원본 {len(posts)}개 글과 별명 {len(aliases)}개를 읽었습니다. 처음부터 평가합니다.")
    for batch_number, batch in enumerate(batches, start=1):
        start = (batch_number - 1) * BATCH_SIZE + 1
        ids = list(range(start, start + len(batch)))
        pending_ids = [post_id for post_id in ids if post_id not in annotations]
        if not pending_ids:
            continue
        pending_posts = [batch[ids.index(post_id)] for post_id in pending_ids]
        print(f"Gemini 평가 중: {batch_number}/{len(batches)} 배치 ({len(pending_posts)}개 제목)")
        annotations.update(evaluate_batch(client, pending_posts, pending_ids, aliases, batch_number))
        save_checkpoint(len(posts), batch_number, annotations)
        if batch_number < len(batches):
            time.sleep(PAUSE_SECONDS)
    write_output(posts, annotations)
    CHECKPOINT_FILE.unlink(missing_ok=True)
    print(f"완료: 결과를 저장했습니다: {OUTPUT_CSV}")


if __name__ == "__main__":
    main()
